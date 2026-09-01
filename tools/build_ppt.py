"""Build Distributed_Traffic_Incident_FA1.pptx (15 slides) with python-pptx.

    python tools/build_ppt.py

Plain academic style: white slides, short text, simple shapes. Screenshots come
from docs/screenshots (real captures). Every statement matches the code.
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs" / "screenshots"
OUT = ROOT / "Distributed_Traffic_Incident_FA1.pptx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
ORANGE = RGBColor(0xE0, 0x7A, 0x2F)
GREY = RGBColor(0x5F, 0x66, 0x73)
LIGHT = RGBColor(0xF2, 0xF3, 0xF5)
LINE = RGBColor(0xD0, 0xD4, 0xDA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xC6, 0x3B, 0x2F)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ helpers
def text(slide, x, y, w, h, lines, size=16, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, after=4):
    """lines: list of str, or list of [(text, {bold,size,color,italic}), ...] runs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        runs = line if isinstance(line, list) else [(line, {})]
        runs = [(r, {}) if isinstance(r, str) else r for r in runs]
        for rtext, o in runs:
            r = p.add_run()
            r.text = rtext
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def rect(slide, x, y, w, h, fill=LIGHT, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def arrow(slide, x1, y1, x2, y2, color=GREY, dashed=False, both=False, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    if dashed:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    if both:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle"}))


def box(slide, x, y, w, h, title, sub=None, fill=WHITE, border=NAVY, size=13):
    rect(slide, x, y, w, h, fill=fill, line=border)
    lines = [[(title, {"bold": True, "size": size})]]
    if sub:
        lines.append([(sub, {"size": size - 3, "color": GREY})])
    text(slide, x, y, w, h, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)


def number(slide, x, y, n, color=ORANGE, d=Inches(0.42)):
    rect(slide, x, y, d, d, fill=color, shape=MSO_SHAPE.OVAL)
    text(slide, x, y, d, d, [[(str(n), {"bold": True, "size": 13, "color": WHITE})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)


def numbered(slide, x, y, w, n, heading, body, size=13):
    number(slide, x, y + Inches(0.02), n)
    text(slide, x + Inches(0.55), y - Inches(0.02), w - Inches(0.55), Inches(1.4),
         [[(heading, {"bold": True, "size": size + 2})], [(body, {"size": size, "color": GREY})]])


def picture(slide, path, x, y, max_w, max_h, caption=None):
    if not path.exists():
        rect(slide, x, y, max_w, max_h, fill=LIGHT, line=LINE)
        return
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    pic = slide.shapes.add_picture(str(path), x + (max_w - w) // 2, y, w, h)
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    if caption:
        text(slide, x, y + h + Inches(0.05), max_w, Inches(0.3), [caption], size=10, color=GREY,
             align=PP_ALIGN.CENTER)


def new_slide(label, title):
    s = prs.slides.add_slide(BLANK)
    text(s, Inches(0.6), Inches(0.35), Inches(8), Inches(0.3),
         [[(label.upper(), {"size": 11, "bold": True, "color": ORANGE})]])
    text(s, Inches(0.6), Inches(0.6), Inches(12), Inches(0.7),
         [[(title, {"size": 28, "bold": True})]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, Inches(0.6), Inches(1.32), Inches(1.2), Inches(0.05), fill=ORANGE)
    text(s, Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.3),
         [[(str(len(prs.slides)), {"size": 10, "color": GREY})]], align=PP_ALIGN.RIGHT)
    return s


# ================================================================= 1 TITLE
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, Inches(13.333), Inches(7.5), fill=NAVY)
text(s, Inches(0.9), Inches(0.9), Inches(8), Inches(0.35),
     [[("FORMATIVE ASSESSMENT 1  ·  DISTRIBUTED SYSTEMS", {"size": 12, "bold": True, "color": ORANGE})]])
text(s, Inches(0.9), Inches(1.6), Inches(11.5), Inches(2.2),
     [[("Distributed Real-Time Traffic Incident", {"size": 40, "bold": True, "color": WHITE})],
      [("Detection, Coordination and Response System", {"size": 40, "bold": True, "color": WHITE})]], spacing=1.0)
text(s, Inches(0.9), Inches(3.75), Inches(10.5), Inches(1.0),
     [[("A Distributed Systems approach to real-time traffic incident communication and "
        "emergency response coordination.", {"size": 17, "color": RGBColor(0xC9, 0xD1, 0xDE)})]])
rect(s, Inches(0.9), Inches(4.95), Inches(11.5), Inches(0.03), fill=ORANGE)
text(s, Inches(0.9), Inches(5.15), Inches(6), Inches(1.6),
     [[("Presented by", {"size": 13, "bold": True, "color": ORANGE})],
      [("[TEAM MEMBER 1]  -  [ROLL NUMBER]", {"size": 15, "color": WHITE})],
      [("[TEAM MEMBER 2]  -  [ROLL NUMBER]", {"size": 15, "color": WHITE})]], after=6)
text(s, Inches(7.2), Inches(5.15), Inches(5.5), Inches(1.6),
     [[("College", {"size": 13, "bold": True, "color": ORANGE})],
      [("[COLLEGE NAME]", {"size": 15, "color": WHITE})],
      [("[DEPARTMENT]  ·  [CLASS / YEAR]", {"size": 13, "color": RGBColor(0xC9, 0xD1, 0xDE)})]], after=6)

# ====================================================== 2 PROBLEM STATEMENT
s = new_slide("Problem Statement", "Traffic Incidents Need Faster Distributed Coordination")
text(s, Inches(0.6), Inches(1.7), Inches(7.2), Inches(2.4),
     [["Accidents, road blockages and signal failures happen without warning. Each traffic control "
       "point keeps its own information, so a neighbouring region often learns about an incident late, "
       "response units are arranged by phone calls, and when one control point goes down the others "
       "lose visibility along with it."]], size=16, spacing=1.25)
text(s, Inches(0.6), Inches(4.15), Inches(7.2), Inches(0.4), [[("Proposed Approach", {"bold": True, "size": 16, "color": ORANGE})]])
text(s, Inches(0.6), Inches(4.55), Inches(7.2), Inches(2.2),
     [["A network of independent regional traffic nodes (Pune, Mumbai, Nashik) that each own their "
       "incident data and talk to each other through RPC, asynchronous messages, live streams and direct "
       "peer-to-peer calls - so incidents are detected, shared and answered in real time, even when a "
       "node fails."]], size=15, color=GREY, spacing=1.25)
problems = [("Isolated Control Points", "Incident and unit data of one region is not visible to the next region."),
            ("Slow Response Coordination", "Borrowing an ambulance or patrol unit from a neighbour is a manual process."),
            ("Single Point of Failure", "If one control point is down, nobody else can record or serve its incidents.")]
for i, (h, b) in enumerate(problems):
    y = Inches(1.7) + i * Inches(1.6)
    rect(s, Inches(8.4), y, Inches(4.4), Inches(1.35), fill=LIGHT)
    text(s, Inches(8.6), y + Inches(0.1), Inches(0.4), Inches(0.5), [[("!", {"bold": True, "size": 22, "color": ORANGE})]])
    text(s, Inches(9.0), y + Inches(0.12), Inches(3.7), Inches(1.2),
         [[(h, {"bold": True, "size": 15})], [(b, {"size": 12.5, "color": GREY})]])

# ============================================================== 3 MOTIVATION
s = new_slide("Motivation", "Why This Problem Matters")
items = [("Time-Critical", "Traffic incidents need a quick response; minutes spent finding a free unit matter."),
         ("Distributed Information", "Pune, Mumbai and Nashik each maintain their own local incident and unit data."),
         ("Node Failure", "One unavailable traffic node should not stop the whole network."),
         ("Regional Coordination", "Nearby traffic nodes should be able to cooperate and share response units.")]
for i, (h, b) in enumerate(items):
    x = Inches(0.8) + (i % 2) * Inches(6.2)
    y = Inches(2.0) + (i // 2) * Inches(2.4)
    numbered(s, x, y, Inches(5.6), i + 1, h, b, size=14)

# ============================================================== 4 OBJECTIVES
s = new_slide("Objectives", "What the Project Sets Out to Do")
objs = ["Detect and report traffic incidents in real time.",
        "Enable communication between geographically distributed traffic nodes.",
        "Coordinate response resources across nodes.",
        "Demonstrate RPC and message-oriented communication.",
        "Demonstrate stream-oriented and P2P communication.",
        "Build a foundation for advanced FA-2 distributed algorithms."]
for i, o in enumerate(objs):
    x = Inches(0.8) + (i % 2) * Inches(6.2)
    y = Inches(1.9) + (i // 2) * Inches(1.6)
    number(s, x, y, i + 1, color=NAVY)
    text(s, x + Inches(0.6), y - Inches(0.05), Inches(5.4), Inches(1.2), [o], size=16, anchor=MSO_ANCHOR.TOP)

# ================================================ 5 DEFINITION / GOALS / TYPES
s = new_slide("Unit 1", "Definition, Goals & Types of Distributed Systems")
text(s, Inches(0.6), Inches(1.65), Inches(12.1), Inches(0.9),
     [[("Definition:  ", {"bold": True, "size": 15}),
       ("A distributed system consists of independent computing nodes that communicate over a network "
        "and cooperate to provide a unified service. Here, three traffic nodes run as separate processes "
        "with separate databases, and the dashboard sees one coordinated network.", {"size": 15})]], spacing=1.2)
text(s, Inches(0.6), Inches(2.75), Inches(3), Inches(0.4), [[("Goals:", {"bold": True, "size": 15})]])
goals = [("Resource Sharing", "One region's response units can be reserved by another region (Pune uses Mumbai's MT-03)."),
         ("Transparency", "The operator reports an incident; which node stores it and lends units stays hidden."),
         ("Scalability", "A new region is one more node process and one registry entry."),
         ("Openness", "Every node follows the same gRPC contract (traffic.proto) and message types."),
         ("Fault Handling", "Failed nodes are detected by heartbeats and requests are redirected.")]
for i, (h, b) in enumerate(goals):
    x = Inches(0.6) + (i % 3) * Inches(4.1)
    y = Inches(3.15) + (i // 3) * Inches(1.25)
    text(s, x, y, Inches(3.9), Inches(1.2),
         [[("✓ " + h, {"bold": True, "size": 14, "color": GREEN})], [(b, {"size": 12.5, "color": GREY})]])
rect(s, Inches(0.6), Inches(5.75), Inches(12.1), Inches(1.15), fill=LIGHT)
text(s, Inches(0.8), Inches(5.8), Inches(11.8), Inches(1.1),
     [[("Types of DS applied:  ", {"bold": True, "size": 14}),
       ("a Distributed Information System (incident, unit and traffic records owned by each regional node) "
        "with Distributed Real-Time / Pervasive characteristics (traffic density streamed every few seconds, "
        "nodes leaving and rejoining the network).", {"size": 14})]], anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)

# ============================================================ 6 ARCHITECTURE
s = new_slide("Architecture", "System Architecture")
L = Inches(0.6)
box(s, L + Inches(1.2), Inches(1.7), Inches(4.8), Inches(0.6), "Traffic Control Dashboard", "browser - HTML / CSS / JS")
arrow(s, L + Inches(3.6), Inches(2.3), L + Inches(3.6), Inches(2.7), both=True)
text(s, L + Inches(3.7), Inches(2.32), Inches(2.5), Inches(0.3), [[("REST + WebSocket", {"size": 10, "color": GREY})]])
box(s, L + Inches(1.2), Inches(2.7), Inches(4.8), Inches(0.7), "Coordination Middleware", "gateway - FastAPI, message broker, failure detector", fill=LIGHT)
xs = [L, L + Inches(2.55), L + Inches(5.1)]
for (nid, port, db), x in zip([("Pune Node", ":5001", "pune.db"), ("Mumbai Node", ":5002", "mumbai.db"), ("Nashik Node", ":5003", "nashik.db")], xs):
    arrow(s, L + Inches(3.6), Inches(3.4), x + Inches(1.05), Inches(4.1))
    box(s, x, Inches(4.1), Inches(2.1), Inches(0.7), nid, "gRPC server " + port, border=GREEN)
    arrow(s, x + Inches(1.05), Inches(4.8), x + Inches(1.05), Inches(5.35), head=False)
    can = s.shapes.add_shape(MSO_SHAPE.CAN, x + Inches(0.6), Inches(5.35), Inches(0.9), Inches(0.7))
    can.fill.solid(); can.fill.fore_color.rgb = LIGHT; can.line.color.rgb = GREY; can.shadow.inherit = False
    text(s, x + Inches(0.3), Inches(5.55), Inches(1.5), Inches(0.3), [[(db, {"size": 10})]], align=PP_ALIGN.CENTER)
text(s, L + Inches(4.9), Inches(3.55), Inches(1.5), Inches(0.3), [[("RPC (gRPC)", {"size": 10, "color": GREY})]])
arrow(s, xs[0] + Inches(2.1), Inches(4.45), xs[1], Inches(4.45), color=ORANGE, dashed=True, both=True)
arrow(s, xs[1] + Inches(2.1), Inches(4.45), xs[2], Inches(4.45), color=ORANGE, dashed=True, both=True)
text(s, xs[0] + Inches(2.1), Inches(4.15), Inches(0.45), Inches(0.3), [[("P2P", {"size": 9, "bold": True, "color": ORANGE})]], align=PP_ALIGN.CENTER)
text(s, xs[1] + Inches(2.1), Inches(4.15), Inches(0.45), Inches(0.3), [[("P2P", {"size": 9, "bold": True, "color": ORANGE})]], align=PP_ALIGN.CENTER)
text(s, L, Inches(6.2), Inches(7.2), Inches(0.7),
     [[("Local SQLite databases - one per node, never shared.   WebRTC: browser tab ↔ browser tab (signalling over the WebSocket).",
        {"size": 10.5, "color": GREY})]])
parts = [("1. Traffic Dashboard", "Operator reports incidents and watches the network live."),
         ("2. Coordination Middleware", "Allocates incident ids, calls nodes by RPC, broadcasts incidents, detects failures and redirects requests."),
         ("3. Regional Traffic Nodes", "Independent processes for Pune, Mumbai and Nashik; each keeps its own incidents, units and traffic status."),
         ("4. P2P Communication", "When local units are insufficient a node asks its nearest peer directly."),
         ("5. Message + Stream Layer", "Broker messages (INCIDENT_CREATED, NODE_FAILED ...) and a WebSocket stream to the dashboard."),
         ("6. Emergency Communication", "WebRTC data channel and optional audio/video between control room and response team.")]
for i, (h, b) in enumerate(parts):
    y = Inches(1.65) + i * Inches(0.88)
    text(s, Inches(8.1), y, Inches(4.8), Inches(0.85),
         [[(h, {"bold": True, "size": 13})], [(b, {"size": 11.5, "color": GREY})]])

# ================================================= 7 DESIGN ISSUES & MIDDLEWARE
s = new_slide("Unit 1", "Design Issues & Middleware")
rows = [("Heterogeneity", "All nodes communicate through one common gRPC interface and one message format."),
        ("Openness", "A new regional node joins by implementing the same contract and registering its id and address."),
        ("Scalability", "Regions are independent processes with their own database; more regions can be added independently."),
        ("Failure Handling", "Heartbeats detect a failed node; its requests are redirected to the nearest available node."),
        ("Concurrency", "Nodes handle requests asynchronously; a unit is reserved with one atomic database update."),
        ("Transparency", "The operator never needs to know which node holds a record or which node lent a unit.")]
tbl = s.shapes.add_table(len(rows) + 1, 2, Inches(0.6), Inches(1.65), Inches(12.1), Inches(3.9)).table
tbl.columns[0].width = Inches(2.6)
tbl.columns[1].width = Inches(9.5)
for c, h in enumerate(("Design Issue", "How Our System Addresses It")):
    cell = tbl.cell(0, c); cell.text = h
    cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
    r = cell.text_frame.paragraphs[0].runs[0]; r.font.bold = True; r.font.size = Pt(14); r.font.color.rgb = WHITE; r.font.name = FONT
for i, (a, b) in enumerate(rows, 1):
    for c, v in enumerate((a, b)):
        cell = tbl.cell(i, c); cell.text = v
        cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
        r = cell.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(13); r.font.name = FONT
        r.font.bold = c == 0; r.font.color.rgb = NAVY if c == 0 else GREY
text(s, Inches(0.6), Inches(5.8), Inches(12.1), Inches(1.1),
     [[("Middleware:  ", {"bold": True, "size": 14, "color": ORANGE}),
       ("the layer between the network and the application that hides the differences between nodes. "
        "In this project the Coordination Service plays that role - it resolves node ids to addresses, makes the "
        "gRPC calls, runs the message broker and applies the failover rule, so the dashboard only calls one API.",
        {"size": 14})]], spacing=1.2)

# ================================= 8 MULTIMEDIA / COMPUTATION / VIRTUALIZATION
s = new_slide("Unit 1", "Distributed Multimedia · Model of Computation · Virtualization")
cols = [("Distributed Multimedia Systems",
         "Transmission of continuous media (audio / video / data) between distributed endpoints in real time.",
         "The Emergency Communication panel opens a WebRTC peer connection between the Traffic Control Center "
         "and the Emergency Response Team (two browser tabs). A data channel carries messages directly "
         "browser-to-browser; audio/video can be added with one button (optional)."),
        ("Model of Distributed Computation",
         "Independent processes with no shared memory and no global clock, interacting only by exchanging messages.",
         "Each traffic node is a separate process with its own database. Nodes interact only through gRPC "
         "calls, broker messages and P2P calls - there is no shared state."),
        ("Role of Virtualization",
         "Running each service in its own isolated container or VM so it can be deployed and restarted independently.",
         "Each regional node can be deployed as an independent Docker container (docker-compose with three node "
         "containers and one gateway). The same code also runs as plain Python processes with run.py.")]
for i, (h, concept, applied) in enumerate(cols):
    x = Inches(0.6) + i * Inches(4.1)
    text(s, x, Inches(1.7), Inches(3.8), Inches(0.8), [[(h, {"bold": True, "size": 16})]])
    text(s, x, Inches(2.45), Inches(3.8), Inches(1.5),
         [[("Concept", {"bold": True, "size": 12, "color": ORANGE})], [(concept, {"size": 13, "color": GREY})]], spacing=1.2)
    rect(s, x, Inches(3.95), Inches(3.8), Inches(0.02), fill=LINE)
    text(s, x, Inches(4.1), Inches(3.8), Inches(2.7),
         [[("Applied in our project", {"bold": True, "size": 12, "color": ORANGE})], [(applied, {"size": 13})]], spacing=1.2)

# ============================================== 9 COMMUNICATION FUNDAMENTALS & RPC
s = new_slide("Unit 2", "Communication Fundamentals & Remote Procedure Call")
text(s, Inches(0.6), Inches(1.65), Inches(12.1), Inches(0.9),
     [[("Fundamentals:  ", {"bold": True, "size": 14}),
       ("distributed nodes communicate by sending messages over the network. Communication can be synchronous "
        "(the caller waits - our gRPC calls) or asynchronous (fire-and-forget - our broker messages).", {"size": 14})]], spacing=1.2)
text(s, Inches(0.6), Inches(2.6), Inches(12.1), Inches(0.5),
     [[("RPC lets the Coordination Service call a procedure on a remote traffic node as if it were local:", {"size": 14, "bold": True})]])
steps = [("Client", "The Coordination Service calls search_response_units(\"Pune\", \"High\") - in code: RequestResponseSupport(INC-1004) - like a normal function."),
         ("Client stub", "The stub generated from traffic.proto packs the parameters (protobuf) and sends them to the Pune node at localhost:5001."),
         ("Remote procedure", "The server stub on the Pune node unpacks the call and runs the local procedure: search own units, then ask peers."),
         ("Response", "The result (incident, assigned units, trace) is returned to the caller - or a gRPC error such as UNAVAILABLE if the node is down.")]
for i, (h, b) in enumerate(steps):
    numbered(s, Inches(0.8), Inches(3.2) + i * Inches(0.85), Inches(7.0), i + 1, h, b, size=12)
rect(s, Inches(8.4), Inches(3.2), Inches(4.4), Inches(3.4), fill=LIGHT)
text(s, Inches(8.6), Inches(3.3), Inches(4.0), Inches(0.4), [[("RPC flow", {"bold": True, "size": 13})]])
for i, (t, sub) in enumerate([("Dashboard", "POST /api/incidents"), ("Coordination Service", "allocates INC-1004"),
                              ("Pune Node  localhost:5001", "records + searches units"), ("Response", "back to the dashboard")]):
    y = Inches(3.7) + i * Inches(0.72)
    box(s, Inches(8.7), y, Inches(3.8), Inches(0.5), t, sub, size=12)
    if i < 3:
        arrow(s, Inches(10.6), y + Inches(0.5), Inches(10.6), y + Inches(0.72))
text(s, Inches(0.8), Inches(6.7), Inches(12), Inches(0.4),
     [[("search_response_units( \"Pune\", \"High\" )   →   RPC   →   Pune Node", {"size": 13, "bold": True, "color": ORANGE})]])

# ========================================== 10 MESSAGE & STREAM COMMUNICATION
s = new_slide("Unit 2", "Message-Oriented & Stream-Oriented Communication")
halves = [("Message-Oriented Communication",
           "Asynchronous, queue-based communication: the sender publishes a message and continues; receivers "
           "pick it up from the queue. Sender and receiver do not wait for each other.",
           "Used for", ["Incident notifications  (INCIDENT_CREATED, INCIDENT_UPDATED)",
                        "Response requests and replies  (RESPONSE_REQUESTED, RESPONSE_AVAILABLE)",
                        "Node failure events  (NODE_FAILED, NODE_RECOVERED)"],
           "In our project: a small publish/subscribe message broker written in Python (queue + worker + "
           "subscribers). Nodes publish, the coordinator and the dashboard subscribe. Broadcasts for a node "
           "that is offline are queued and delivered when it recovers."),
          ("Stream-Oriented Communication",
           "A continuous, time-ordered flow of data where timing matters - receivers get updates as they "
           "happen instead of asking again and again.",
           "Used for", ["Traffic density of each region", "Node heartbeat / health status",
                        "Live incident updates and the communication log"],
           "In our project: a WebSocket (/ws) from the gateway to the dashboard. Every broker message and a "
           "status snapshot every 2 seconds are pushed to the browser, so the page never needs a refresh.")]
for i, (h, concept, u, uses, impl) in enumerate(halves):
    x = Inches(0.6) + i * Inches(6.3)
    text(s, x, Inches(1.7), Inches(5.8), Inches(0.5), [[(h, {"bold": True, "size": 17})]])
    text(s, x, Inches(2.2), Inches(5.8), Inches(1.2), [concept], size=13.5, color=GREY, spacing=1.2)
    text(s, x, Inches(3.4), Inches(5.8), Inches(0.35), [[(u, {"bold": True, "size": 13, "color": ORANGE})]])
    text(s, x, Inches(3.75), Inches(5.8), Inches(1.3), [["•  " + item] for item in uses], size=13, after=3)
    rect(s, x, Inches(5.05), Inches(5.8), Inches(1.8), fill=LIGHT)
    text(s, x + Inches(0.15), Inches(5.1), Inches(5.5), Inches(1.7), [impl], size=12.5, anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)

# ===================================================== 11 P2P & WEBRTC
s = new_slide("Unit 2", "Peer-to-Peer Messaging & WebRTC")
# left
x = Inches(0.6)
text(s, x, Inches(1.7), Inches(5.8), Inches(0.5), [[("P2P Messaging", {"bold": True, "size": 17})]])
text(s, x, Inches(2.2), Inches(5.8), Inches(1.0),
     ["Nodes communicate directly with one another, without going through a central server or a shared "
      "database - so there is no single point of failure for that exchange."], size=13.5, color=GREY, spacing=1.2)
box(s, x, Inches(3.35), Inches(2.2), Inches(0.7), "Pune Node", "localhost:5001", border=GREEN)
box(s, x + Inches(3.6), Inches(3.35), Inches(2.2), Inches(0.7), "Mumbai Node", "localhost:5002", border=GREEN)
arrow(s, x + Inches(2.2), Inches(3.55), x + Inches(3.6), Inches(3.55), color=ORANGE)
arrow(s, x + Inches(3.6), Inches(3.85), x + Inches(2.2), Inches(3.85), color=ORANGE)
text(s, x + Inches(2.2), Inches(3.1), Inches(1.4), Inches(0.3), [[("request units", {"size": 9.5, "color": ORANGE})]], align=PP_ALIGN.CENTER)
text(s, x + Inches(2.2), Inches(4.05), Inches(1.4), Inches(0.3), [[("MT-03 available", {"size": 9.5, "color": ORANGE})]], align=PP_ALIGN.CENTER)
text(s, x, Inches(4.55), Inches(5.8), Inches(2.4),
     [[("Applied in our project", {"bold": True, "size": 13, "color": ORANGE})],
      ["When Pune's own units are not enough, the Pune node itself calls the Mumbai node's "
       "SearchResponseUnits and ReserveUnit procedures - nearest peer first, next peer if that one is down or "
       "empty. The exchange is shown live in the dashboard log as  PUNE → MUMBAI  /  MUMBAI → PUNE."]],
     size=13, spacing=1.2)
# right
x = Inches(6.9)
text(s, x, Inches(1.7), Inches(5.8), Inches(0.5), [[("WebRTC", {"bold": True, "size": 17})]])
text(s, x, Inches(2.2), Inches(5.8), Inches(1.0),
     ["An open browser standard for direct, low-latency peer connections carrying audio, video and data. "
      "A server is only needed to exchange the connection details (signalling)."], size=13.5, color=GREY, spacing=1.2)
box(s, x, Inches(3.35), Inches(2.3), Inches(0.7), "Traffic Control Center", "browser tab 1", size=12)
box(s, x + Inches(3.5), Inches(3.35), Inches(2.3), Inches(0.7), "Emergency Response Team", "browser tab 2", size=12)
arrow(s, x + Inches(2.3), Inches(3.7), x + Inches(3.5), Inches(3.7), color=ORANGE, both=True)
text(s, x + Inches(2.3), Inches(3.1), Inches(1.2), Inches(0.3), [[("data channel", {"size": 9.5, "color": ORANGE})]], align=PP_ALIGN.CENTER)
text(s, x + Inches(2.3), Inches(4.05), Inches(1.2), Inches(0.3), [[("+ audio / video", {"size": 9.5, "color": ORANGE})]], align=PP_ALIGN.CENTER)
text(s, x, Inches(4.55), Inches(5.8), Inches(2.4),
     [[("Applied in our project", {"bold": True, "size": 13, "color": ORANGE})],
      ["Two dashboard tabs connect with \"Connect Peer\": the offer, answer and ICE candidates travel over our "
       "WebSocket, then messages flow directly between the browsers through a data channel. "
       "\"Start Audio/Video\" adds camera and microphone (optional, needs a device)."]],
     size=13, spacing=1.2)

# ============================================= 12 NAMES / ADDRESSES + FAULT TOLERANCE
s = new_slide("Unit 2", "Names, Identifiers & Addresses  ·  Fault Tolerance")
x = Inches(0.6)
text(s, x, Inches(1.7), Inches(5.8), Inches(0.5), [[("Names, Identifiers & Addresses", {"bold": True, "size": 17})]])
text(s, x, Inches(2.2), Inches(5.8), Inches(0.9),
     ["A name or identifier says what an entity is; an address says where to find it. Good distributed "
      "design keeps the two separate."], size=13.5, color=GREY, spacing=1.2)
reg = [("Node ID", "Node Name", "Address"), ("PUNE-TR-01", "Pune Traffic Node", "localhost:5001"),
       ("MUM-TR-01", "Mumbai Traffic Node", "localhost:5002"), ("NSK-TR-01", "Nashik Traffic Node", "localhost:5003")]
t2 = s.shapes.add_table(4, 3, x, Inches(3.2), Inches(5.8), Inches(1.6)).table
for j, w in enumerate((1.7, 2.3, 1.8)):
    t2.columns[j].width = Inches(w)
for i, row in enumerate(reg):
    for j, v in enumerate(row):
        cell = t2.cell(i, j); cell.text = v
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if i == 0 else (WHITE if i % 2 else LIGHT)
        r = cell.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(12); r.font.name = FONT
        r.font.bold = i == 0; r.font.color.rgb = WHITE if i == 0 else NAVY
text(s, x, Inches(4.95), Inches(5.8), Inches(2.0),
     [[("Applied in our project", {"bold": True, "size": 13, "color": ORANGE})],
      ["Every message and record refers to a node by its id, never by host:port; a registry resolves the id "
       "to an address (in Docker the same ids map to different addresses). Incidents get unique ids "
       "INC-1001, INC-1002 ... and units PT-01, MT-03, NT-02 carry their owner node."]], size=13, spacing=1.2)
x = Inches(6.9)
text(s, x, Inches(1.7), Inches(5.8), Inches(0.5), [[("Fault Tolerance", {"bold": True, "size": 17})]])
text(s, x, Inches(2.2), Inches(5.8), Inches(0.9),
     ["The ability of the system to keep working - possibly at reduced capacity - when some of its "
      "components fail."], size=13.5, color=GREY, spacing=1.2)
chain = [("Pune OFFLINE", "heartbeat fails → NODE_FAILED", RED), ("Request redirected", "to the nearest online node", ORANGE),
         ("Mumbai responds", "records the incident, borrows units from Nashik by P2P", GREEN)]
for i, (h, b, c) in enumerate(chain):
    y = Inches(3.2) + i * Inches(0.85)
    box(s, x + Inches(0.6), y, Inches(4.6), Inches(0.62), h, b, border=c, size=13)
    if i < 2:
        arrow(s, x + Inches(2.9), y + Inches(0.62), x + Inches(2.9), y + Inches(0.85))
text(s, x, Inches(5.8), Inches(5.8), Inches(1.2),
     [[("Simulation note:  ", {"bold": True, "size": 12, "color": RED}),
       ("\"Fail Pune Node\" makes the node refuse every RPC with gRPC status UNAVAILABLE, which is what a crashed "
        "node looks like to callers. Killing the process is detected the same way. \"Recover\" brings it back "
        "and it receives the broadcasts it missed.", {"size": 12, "color": GREY})]], spacing=1.15)

# ============================================================ 13 DEMONSTRATION
s = new_slide("Demonstration", "Working Demo Flow")
demo = [("Report accident in Pune", "Accident · Pune · High from the dashboard form"),
        ("Pune node receives incident", "coordinator allocates INC-1004, RPC CreateIncident → stored in pune.db"),
        ("Incident is broadcast", "INCIDENT_CREATED → Mumbai and Nashik store a copy"),
        ("Pune searches locally", "High needs 2 units, Pune has only PT-02 free"),
        ("Pune requests another node", "P2P SearchResponseUnits to Mumbai"),
        ("Mumbai responds", "MT-03 available → reserved; status UNITS ASSIGNED"),
        ("Dashboard updates live", "table, units and log refresh over the WebSocket"),
        ("Simulate Pune failure", "\"Fail Pune Node\" → heartbeat fails → OFFLINE"),
        ("Request continues", "new Pune incident redirected to Mumbai, units borrowed from Nashik")]
for i, (h, b) in enumerate(demo):
    y = Inches(1.6) + i * Inches(0.6)
    number(s, Inches(0.7), y, i + 1, color=NAVY if i < 7 else RED, d=Inches(0.36))
    text(s, Inches(1.2), y - Inches(0.04), Inches(5.3), Inches(0.55),
         [[(h + "   ", {"bold": True, "size": 12.5}), (b, {"size": 11, "color": GREY})]], anchor=MSO_ANCHOR.MIDDLE, after=0)
picture(s, SHOTS / "03_incident_creation.png", Inches(6.8), Inches(1.6), Inches(3.0), Inches(2.6), "Steps 1-6: request trace")
picture(s, SHOTS / "05c_failover_trace.png", Inches(9.95), Inches(1.6), Inches(3.0), Inches(2.6), "Steps 8-9: redirected after failure")
picture(s, SHOTS / "05d_failover_log.png", Inches(6.8), Inches(4.5), Inches(3.0), Inches(2.4), "Communication log during failover")
picture(s, SHOTS / "05b_network_status_offline.png", Inches(9.95), Inches(4.5), Inches(3.0), Inches(2.4), "Network status: Pune OFFLINE")

# ========================================================== 14 IMPLEMENTATION
s = new_slide("Implementation", "Technology Stack")
stack = [("Frontend", "HTML + CSS + JavaScript", "Single dashboard page, network diagram, live tables and log"),
         ("Backend", "Python + FastAPI", "REST API, coordination middleware, message broker, WebSocket"),
         ("RPC", "Python gRPC", "traffic.proto contract; each node is a gRPC server"),
         ("Real-Time Stream", "WebSocket", "Pushes events and status snapshots to the dashboard"),
         ("Messaging", "Lightweight Python message layer", "Publish/subscribe broker (RabbitMQ not used)"),
         ("P2P", "Python node-to-node gRPC", "Direct SearchResponseUnits / ReserveUnit between nodes"),
         ("WebRTC", "Browser WebRTC", "Data channel + optional audio/video between two tabs"),
         ("Database", "SQLite", "One database file per node: pune.db, mumbai.db, nashik.db"),
         ("Virtualization", "Docker", "One container per node + gateway (docker-compose)")]
for i, (h, tech, b) in enumerate(stack):
    x = Inches(0.6) + (i % 3) * Inches(4.1)
    y = Inches(1.7) + (i // 3) * Inches(1.7)
    rect(s, x, y, Inches(3.85), Inches(1.5), fill=LIGHT)
    text(s, x + Inches(0.2), y + Inches(0.12), Inches(3.5), Inches(1.3),
         [[(h, {"bold": True, "size": 11.5, "color": ORANGE})], [(tech, {"bold": True, "size": 15})],
          [(b, {"size": 12, "color": GREY})]], after=2)
text(s, Inches(0.6), Inches(6.85), Inches(12), Inches(0.3),
     [[("All of the above is implemented and runs on one laptop:  pip install -r requirements.txt  →  python run.py", {"size": 11.5, "color": GREY})]])

# ======================================================== 15 FA-2 ROADMAP + THANKS
s = new_slide("Future Scope", "FA-2 Extension Roadmap")
road = [("Mutual Exclusion", "Avoid conflicting allocation of the same response unit by two nodes."),
        ("Clock Synchronization", "Order events coming from different traffic nodes consistently."),
        ("Election Algorithms", "Select the coordinator dynamically instead of a fixed gateway."),
        ("Replication & Recovery", "Keep a region's data available while its node is down."),
        ("Distributed Consistency", "Keep incident copies synchronized across nodes.")]
for i, (h, b) in enumerate(road):
    x = Inches(0.8) + (i % 3) * Inches(4.1)
    y = Inches(1.8) + (i // 3) * Inches(1.7)
    numbered(s, x, y, Inches(3.7), i + 1, h, b, size=12.5)
text(s, Inches(0.6), Inches(5.2), Inches(12), Inches(0.4),
     [[("These are planned extensions - not part of the FA-1 implementation.", {"size": 12, "italic": True, "color": GREY})]])
rect(s, Inches(0.6), Inches(5.75), Inches(12.1), Inches(1.2), fill=NAVY)
text(s, Inches(0.6), Inches(5.75), Inches(12.1), Inches(1.2),
     [[("Thank You", {"bold": True, "size": 32, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("saved", OUT, "-", len(prs.slides), "slides")
