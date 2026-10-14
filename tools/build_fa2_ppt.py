"""Build Distributed_Traffic_Incident_FA2.pptx (16 slides) with python-pptx.

    python tools/capture_fa2_screenshots.py   (system running; real captures -> docs/screenshots/fa2/)
    python tools/build_fa2_ppt.py

Same plain academic style and helper functions as tools/build_ppt.py (FA-1), which is left
unchanged; the FA-1 presentation is not overwritten. Every statement matches the FA-2 code.
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs" / "screenshots" / "fa2"
OUT = ROOT / "Distributed_Traffic_Incident_FA2.pptx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
ORANGE = RGBColor(0xE0, 0x7A, 0x2F)
GREY = RGBColor(0x5F, 0x66, 0x73)
LIGHT = RGBColor(0xF2, 0xF3, 0xF5)
LINE = RGBColor(0xD0, 0xD4, 0xDA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xC6, 0x3B, 0x2F)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ helpers
def text(slide, x, y, w, h, lines, size=16, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, after=4):
    """lines: list of str, or list of [(text, {bold,size,color,italic}), ...] runs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        runs = line if isinstance(line, list) else [(line, {})]
        runs = [(r, {}) if isinstance(r, str) else r for r in runs]
        for rtext, o in runs:
            r = p.add_run()
            r.text = rtext
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def rect(slide, x, y, w, h, fill=LIGHT, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def arrow(slide, x1, y1, x2, y2, color=GREY, dashed=False, both=False, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    if dashed:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    if both:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle"}))


def box(slide, x, y, w, h, title, sub=None, fill=WHITE, border=NAVY, size=13):
    rect(slide, x, y, w, h, fill=fill, line=border)
    lines = [[(title, {"bold": True, "size": size})]]
    if sub:
        lines.append([(sub, {"size": size - 3, "color": GREY})])
    text(slide, x, y, w, h, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)


def number(slide, x, y, n, color=ORANGE, d=Inches(0.42)):
    rect(slide, x, y, d, d, fill=color, shape=MSO_SHAPE.OVAL)
    text(slide, x, y, d, d, [[(str(n), {"bold": True, "size": 13, "color": WHITE})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)


def numbered(slide, x, y, w, n, heading, body, size=13):
    number(slide, x, y + Inches(0.02), n)
    text(slide, x + Inches(0.55), y - Inches(0.02), w - Inches(0.55), Inches(1.4),
         [[(heading, {"bold": True, "size": size + 2})], [(body, {"size": size, "color": GREY})]])


def picture(slide, path, x, y, max_w, max_h, caption=None):
    if not path.exists():
        rect(slide, x, y, max_w, max_h, fill=LIGHT, line=LINE)
        return
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    pic = slide.shapes.add_picture(str(path), x + (max_w - w) // 2, y, w, h)
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    if caption:
        text(slide, x, y + h + Inches(0.05), max_w, Inches(0.3), [caption], size=10, color=GREY,
             align=PP_ALIGN.CENTER)


def new_slide(label, title):
    s = prs.slides.add_slide(BLANK)
    text(s, Inches(0.6), Inches(0.35), Inches(8), Inches(0.3),
         [[(label.upper(), {"size": 11, "bold": True, "color": ORANGE})]])
    text(s, Inches(0.6), Inches(0.6), Inches(12), Inches(0.7),
         [[(title, {"size": 28, "bold": True})]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, Inches(0.6), Inches(1.32), Inches(1.2), Inches(0.05), fill=ORANGE)
    text(s, Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.3),
         [[(str(len(prs.slides)), {"size": 10, "color": GREY})]], align=PP_ALIGN.RIGHT)
    return s


def table(slide, x, y, w, rows, widths, size=12, row_h=0.42):
    """Header row navy, body rows alternating - same look as the FA-1 design-issues table."""
    t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows))).table
    for j, wd in enumerate(widths):
        t.columns[j].width = Inches(wd)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.cell(i, j)
            cell.text = str(v)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (WHITE if i % 2 else LIGHT)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
                    r.font.name = FONT
                    r.font.bold = i == 0 or j == 0
                    r.font.color.rgb = WHITE if i == 0 else (NAVY if j == 0 else GREY)
    return t


def concept(slide, x, y, w, h, heading, body, applied):
    text(slide, x, y, w, Inches(0.4), [[(heading, {"bold": True, "size": 17})]])
    text(slide, x, y + Inches(0.4), w, Inches(0.75), [body], size=12, color=GREY, spacing=1.1)
    rect(slide, x, y + Inches(1.15), w, h - Inches(1.15), fill=LIGHT)
    text(slide, x + Inches(0.15), y + Inches(1.18), w - Inches(0.3), h - Inches(1.2),
         [[("Applied in our project", {"bold": True, "size": 12, "color": ORANGE})]] + [[(a, {"size": 12})] for a in applied],
         spacing=1.12, after=3)


# ================================================================= 1 TITLE
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, Inches(13.333), Inches(7.5), fill=NAVY)
text(s, Inches(0.9), Inches(0.9), Inches(10), Inches(0.35),
     [[("FORMATIVE ASSESSMENT 2  ·  DISTRIBUTED SYSTEMS  ·  UNIT III & UNIT IV", {"size": 12, "bold": True, "color": ORANGE})]])
text(s, Inches(0.9), Inches(1.6), Inches(11.5), Inches(2.2),
     [[("Distributed Real-Time Traffic Incident", {"size": 40, "bold": True, "color": WHITE})],
      [("Detection, Coordination and Response System", {"size": 40, "bold": True, "color": WHITE})]], spacing=1.0)
text(s, Inches(0.9), Inches(3.75), Inches(11), Inches(1.0),
     [[("FA-2: synchronization (clocks, logical time, global state, election, mutual exclusion, beacons) and "
        "emerging paradigms (objects, web, DFS, serverless, ledger) added to the FA-1 system.",
        {"size": 17, "color": RGBColor(0xC9, 0xD1, 0xDE)})]])
rect(s, Inches(0.9), Inches(4.95), Inches(11.5), Inches(0.03), fill=ORANGE)
text(s, Inches(0.9), Inches(5.15), Inches(6), Inches(1.6),
     [[("Presented by", {"size": 13, "bold": True, "color": ORANGE})],
      [("[TEAM MEMBER 1]  -  [ROLL NUMBER]", {"size": 15, "color": WHITE})],
      [("[TEAM MEMBER 2]  -  [ROLL NUMBER]", {"size": 15, "color": WHITE})]], after=6)
text(s, Inches(7.2), Inches(5.15), Inches(5.5), Inches(1.6),
     [[("College", {"size": 13, "bold": True, "color": ORANGE})],
      [("[COLLEGE NAME]", {"size": 15, "color": WHITE})],
      [("[DEPARTMENT]  ·  [CLASS / YEAR]", {"size": 13, "color": RGBColor(0xC9, 0xD1, 0xDE)})]], after=6)

# ============================================================ 2 FA-2 SCOPE
s = new_slide("FA-2 Overview", "From FA-1 to FA-2: Extending, Not Rebuilding")
text(s, Inches(0.6), Inches(1.65), Inches(12.1), Inches(0.9),
     [[("FA-1 gave us three independent regional nodes (Pune, Mumbai, Nashik) that communicate by RPC, messages, "
        "streams and P2P. FA-2 adds the algorithms that make such nodes ", {"size": 15}),
       ("agree", {"size": 15, "bold": True}), (" - on time, order, a leader and access to shared resources - and "
        "the newer ways of building distributed applications.", {"size": 15})]], spacing=1.2)
units = [("Unit III — Synchronization", ["Clock synchronization: Cristian (+NTP), Berkeley", "Logical clocks: Lamport's algorithm, vector clocks",
                                         "Causal ordering of broadcasts", "Global state: Chandy–Lamport snapshot",
                                         "Election: Bully and Ring", "Mutual exclusion: Lamport, Ricart–Agrawala",
                                         "Beacon protocol", "Self-study: Lodha–Kshemkalyani, Knapp's classification"]),
         ("Unit IV — Emerging Paradigms", ["Distributed object-based system (proxy, skeleton, naming)",
                                          "Distributed web-based system (HTTP caching, ETag/304)",
                                          "Distributed file system (chunks, replicas, checksums)",
                                          "Serverless architecture (local FaaS)", "Self-study: blockchain / distributed ledger",
                                          "Case studies: Megaport, Cloudflare, AWS, Hadoop, Kubernetes"])]
for i, (h, items) in enumerate(units):
    x = Inches(0.6) + i * Inches(6.2)
    rect(s, x, Inches(2.7), Inches(5.9), Inches(3.6), fill=LIGHT)
    text(s, x + Inches(0.2), Inches(2.8), Inches(5.5), Inches(3.4),
         [[(h, {"bold": True, "size": 16, "color": ORANGE})]] + [["•  " + it] for it in items], size=13.5, after=3)
text(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.5),
     [[("Rule followed: FA-1 code, APIs and UI kept intact - FA-2 is a new package, a second gRPC service and a new "
        "dashboard section. FA-1's 40-assertion test still passes.", {"size": 12.5, "color": GREY, "italic": True})]])

# ============================================================ 3 ARCHITECTURE
s = new_slide("Architecture", "How FA-2 Plugs Into the FA-1 System")
L = Inches(0.6)
box(s, L + Inches(1.2), Inches(1.65), Inches(4.8), Inches(0.6), "Dashboard", "FA-1 panels + FA-2 section (fa2.js)")
arrow(s, L + Inches(3.6), Inches(2.25), L + Inches(3.6), Inches(2.6), both=True)
box(s, L + Inches(0.6), Inches(2.6), Inches(6.0), Inches(0.9), "Gateway (FastAPI)",
    "FA-1 coordinator · FA-2: /api/fa2, DFS master, naming, FaaS, ledger service", fill=LIGHT)
xs = [L, L + Inches(2.55), L + Inches(5.1)]
for (nid, port), x in zip([("Pune Node", ":5001"), ("Mumbai Node", ":5002"), ("Nashik Node", ":5003")], xs):
    arrow(s, L + Inches(3.6), Inches(3.5), x + Inches(1.05), Inches(4.2))
    box(s, x, Inches(4.2), Inches(2.1), Inches(0.8), nid, "TrafficNode + SyncNode " + port, border=GREEN, size=12)
    text(s, x, Inches(5.1), Inches(2.1), Inches(0.9),
         [[("FA-2 runtime", {"bold": True, "size": 10})], [("clocks · inbox · algorithms", {"size": 9.5, "color": GREY})],
          [("*.db + *_ledger.db + dfs/", {"size": 9.5, "color": GREY})]], align=PP_ALIGN.CENTER, after=0)
arrow(s, xs[0] + Inches(2.1), Inches(4.6), xs[1], Inches(4.6), color=ORANGE, dashed=True, both=True)
arrow(s, xs[1] + Inches(2.1), Inches(4.6), xs[2], Inches(4.6), color=ORANGE, dashed=True, both=True)
text(s, L, Inches(6.1), Inches(7.2), Inches(0.8),
     [[("Orange: SyncNode.Deliver - FIFO channels between all pairs of nodes; every message carries the sender's "
        "Lamport and vector clock.", {"size": 11, "color": GREY})]])
parts = [("Second gRPC service, same port", "proto/fa2_sync.proto: Deliver, GetTime, Command, Invoke, StoreChunk, ReadChunk. FA-1's traffic.proto is unchanged."),
         ("Message-passing algorithms", "Each algorithm is written as 'on receiving X do ...' handlers that run one at a time (atomic)."),
         ("FIFO channels", "One send queue per peer keeps send order - needed by Lamport ME, Lodha–Kshemkalyani and Chandy–Lamport."),
         ("Same failure model as FA-1", "'Fail Node' makes a node refuse every FA-2 RPC too - peers see a crash."),
         ("Gateway triggers, nodes decide", "The dashboard starts an election or a CS request; the nodes exchange the messages."),
         ("Integration", "Every FA-1 incident create/update is timestamped and causally broadcast; FA-2 events appear in the FA-1 log.")]
for i, (h, b) in enumerate(parts):
    y = Inches(1.65) + i * Inches(0.88)
    text(s, Inches(8.0), y, Inches(4.9), Inches(0.85), [[(h, {"bold": True, "size": 13})], [(b, {"size": 11.5, "color": GREY})]])

# ======================================================== 4 CLOCK SYNCHRONIZATION
s = new_slide("Unit III", "Clock Synchronization: Cristian and Berkeley")
text(s, Inches(0.6), Inches(1.6), Inches(6.2), Inches(1.0),
     ["There is no global clock: each node's clock has an offset and drifts. Each node runs a software clock with a "
      "deliberate offset and drift (simulated - one laptop has one OS clock); the algorithms that correct it are real."],
     size=13, color=GREY, spacing=1.15)
numbered(s, Inches(0.6), Inches(2.7), Inches(6.2), 1, "Cristian's algorithm (external)",
         "Ask the time server (gateway = reference clock): new time = t3 + RTT/2. Four samples, the smallest RTT is used. "
         "NTP offset θ and delay δ are computed from the same four timestamps.", size=12)
numbered(s, Inches(0.6), Inches(3.95), Inches(6.2), 2, "Berkeley algorithm (internal)",
         "The elected leader is the time daemon: polls every clock (GetTime), removes outliers, averages, and sends each "
         "node its own adjustment (CLOCK_ADJUST).", size=12)
rect(s, Inches(0.6), Inches(5.3), Inches(6.2), Inches(1.5), fill=LIGHT)
text(s, Inches(0.8), Inches(5.35), Inches(5.9), Inches(1.4),
     [[("Measured", {"bold": True, "size": 13, "color": ORANGE})],
      ["Cristian: clock errors of 0.8-2.1 s reduced to < 1 ms (first sample: ~265 ms RTT of connection set-up - discarded)."],
      ["Berkeley: spread between the three clocks 2.1 s → about 1 ms."]], size=12, spacing=1.1, after=2)
picture(s, SHOTS / "02a_clocks_skewed.png", Inches(7.1), Inches(1.6), Inches(5.7), Inches(2.4), "Before: skewed node clocks")
picture(s, SHOTS / "02c_berkeley.png", Inches(7.1), Inches(4.35), Inches(5.7), Inches(2.5), "After a Berkeley round led by the leader")

# ====================================================== 5 LOGICAL CLOCKS + CAUSAL
s = new_slide("Unit III", "Logical Clocks: Lamport, Vector and Causal Ordering")
text(s, Inches(0.6), Inches(1.6), Inches(5.6), Inches(3.4),
     [[("Lamport's algorithm", {"bold": True, "size": 14})],
      [("L = L + 1 before every event;  on receive L = max(L, L_msg) + 1.  a → b ⇒ L(a) < L(b).", {"size": 12, "color": GREY})],
      [("Vector algorithm", {"bold": True, "size": 14})],
      [("V[i] += 1 per event;  on receive V = max(V, V_msg), V[i] += 1.  a → b ⇔ V(a) < V(b); incomparable = concurrent.", {"size": 12, "color": GREY})],
      [("Causal broadcast (Birman–Schiper–Stephenson)", {"bold": True, "size": 14})],
      [("Deliver a message from j only if it is j's next one and everything j had seen is already delivered; "
        "otherwise buffer it.", {"size": 12, "color": GREY})]], spacing=1.1, after=4)
rect(s, Inches(0.6), Inches(4.75), Inches(5.6), Inches(2.1), fill=LIGHT)
text(s, Inches(0.8), Inches(4.8), Inches(5.3), Inches(2.0),
     [[("Applied in our project", {"bold": True, "size": 12, "color": ORANGE})],
      ["Every FA-2 message carries both clocks; the gateway merges the three event logs and checks every send/receive "
       "pair: 0 Lamport and 0 vector violations, while the skewed physical times put the receive before the send in "
       "many pairs."],
      ["Demo: Pune→Nashik slowed 3 s; Nashik receives Mumbai's ACK first, buffers it, delivers ALERT then ACK."]],
     size=11.5, spacing=1.1, after=3)
picture(s, SHOTS / "04_logical_causal.png", Inches(6.45), Inches(1.6), Inches(6.4), Inches(5.0),
        "Real event log ordered by (Lamport, node) · Nashik receive vs delivery order")

# ============================================================ 6 GLOBAL STATE
s = new_slide("Unit III", "Global State: Chandy–Lamport Snapshot")
text(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(1.0),
     ["A global state = the local state of every node + the messages still travelling in the channels. It must be a "
      "consistent cut: no message received in the cut without its send also being in the cut."], size=13, color=GREY, spacing=1.15)
steps = [("Initiator", "records its state, sends MARKER on every channel, starts recording incoming channels."),
         ("First MARKER at a node", "record own state, that channel is empty, forward MARKERs, record the other channels."),
         ("Later MARKER", "stop recording that channel - what was recorded is 'in transit'."),
         ("FIFO channels", "guarantee a marker separates messages sent before / after the cut.")]
for i, (h, b) in enumerate(steps):
    numbered(s, Inches(0.6), Inches(2.7) + i * Inches(0.8), Inches(6.0), i + 1, h, b, size=12)
rect(s, Inches(0.6), Inches(6.0), Inches(6.0), Inches(0.9), fill=LIGHT)
text(s, Inches(0.8), Inches(6.0), Inches(5.7), Inches(0.9),
     [[("Application: ", {"bold": True, "size": 12, "color": ORANGE}),
       ("30 standby patrol vehicles are lent between regions (0.2-1.2 s in transit). Snapshot: local + in transit = 30 every "
        "time; a naive reading of the three counters does not add up.", {"size": 12})]], anchor=MSO_ANCHOR.MIDDLE, spacing=1.1)
picture(s, SHOTS / "08_snapshot.png", Inches(6.9), Inches(1.6), Inches(5.9), Inches(4.5), "Snapshot recorded by the Nashik leader")

# ======================================================= 7 ELECTION + BEACONS
s = new_slide("Unit III", "Election Algorithms and the Beacon Protocol")
text(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(2.6),
     [[("Bully", {"bold": True, "size": 14})],
      [("Priority Pune 1 < Mumbai 2 < Nashik 3. ELECTION to higher nodes; OK stops lower ones; the highest live node "
        "sends COORDINATOR. Timeouts detect dead nodes.", {"size": 12, "color": GREY})],
      [("Ring", {"bold": True, "size": 14})],
      [("RING_ELECTION collects live ids around Pune→Mumbai→Nashik (dead successor skipped), then RING_COORDINATOR "
        "circulates once: 2n = 6 messages.", {"size": 12, "color": GREY})],
      [("Beacon protocol", {"bold": True, "size": 14})],
      [("Every node sends a timestamped beacon to every peer each second: liveness (no beacon for 3.5 s → SUSPECTED) and "
        "clock-offset estimates (sender→receiver and reference-broadcast).", {"size": 12, "color": GREY})]], spacing=1.1, after=3)
rect(s, Inches(0.6), Inches(4.75), Inches(6.0), Inches(2.1), fill=LIGHT)
text(s, Inches(0.8), Inches(4.8), Inches(5.7), Inches(2.0),
     [[("Demo with FA-1's Fail button", {"bold": True, "size": 12, "color": ORANGE})],
      ["1. Fail Nashik (the leader) → Pune and Mumbai stop hearing its beacons → SUSPECTED."],
      ["2. Suspected leader triggers Bully → Mumbai becomes coordinator (Berkeley daemon, ledger proposer)."],
      ["3. Recover Nashik → it starts an election and bullies its way back."]], size=11.5, spacing=1.1, after=2)
picture(s, SHOTS / "06c_leader_failed_election.png", Inches(6.9), Inches(1.6), Inches(5.9), Inches(3.3),
        "Bully triggered by missing beacons → Mumbai elected")
picture(s, SHOTS / "06b_leader_failed_beacons.png", Inches(6.9), Inches(5.25), Inches(5.9), Inches(1.6),
        "Beacon table: Nashik SUSPECTED")

# ======================================================= 8 MUTUAL EXCLUSION
s = new_slide("Unit III", "Distributed Mutual Exclusion on a Shared Resource")
text(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(0.8),
     [[("Shared resource: ", {"bold": True, "size": 13}),
       ("the regional air-ambulance register, a file in our DFS. The critical section reads it, works 0.6 s, appends a "
        "line and writes it back - if two nodes overlap, an update is lost.", {"size": 13})]], spacing=1.15)
table(s, Inches(0.6), Inches(2.5), Inches(6.6),
      [("Algorithm", "Messages / CS", "Measured (3 concurrent)", "Lost updates"),
       ("Lamport", "3(N-1)", "18", "0"), ("Ricart–Agrawala", "2(N-1)", "12", "0"),
       ("Lodha–Kshemkalyani", "N-1 … 2(N-1)", "8", "0"), ("None", "0", "0", "2  (overlap!)")],
      (2.0, 1.4, 1.9, 1.3), size=12)
text(s, Inches(7.5), Inches(2.5), Inches(5.3), Inches(2.2),
     [["•  Lamport: request queue + REPLY + RELEASE; enter when own request is first and a later-stamped message came from everyone."],
      ["•  Ricart–Agrawala: reply at once unless own request has priority - then defer the reply until exit."],
      ["•  All three grant the CS in timestamp order (fair); the 'None' run proves why an algorithm is needed."]],
     size=12, spacing=1.1, after=3)
picture(s, SHOTS / "07b_mutex_compare_none.png", Inches(0.6), Inches(4.75), Inches(12.1), Inches(1.95),
        "Runs table + timeline of the unsafe 'None' run")

# ======================================================= 9 SELF STUDY LK
s = new_slide("Self-Study", "Lodha & Kshemkalyani's Fair Mutual Exclusion Algorithm")
text(s, Inches(0.6), Inches(1.55), Inches(12.1), Inches(0.6),
     [[("IEEE TPDS 11(6), 2000 - implemented exactly as Fig. 2 of the paper (backend/fa2/mutex.py, class LodhaKshemkalyani).",
        {"size": 12, "italic": True, "color": GREY})]])
lk = [("Request", "SN = highest seen + 1; LRQ = [own request]; REQUEST to all; RV = 0 except own."),
      ("On REQUEST (while requesting)", "RV[j] = 0 → concurrent: insert in LRQ, RV[j] = 1 - the request itself is the reply. RV[j] = 1 → defer."),
      ("On REQUEST (not requesting)", "send REPLY(last satisfied request)."),
      ("On REPLY / FLUSH(Rj)", "RV[j] = 1; remove every LRQ entry with priority ≥ Rj (a collective reply)."),
      ("Enter CS", "all RV = 1 and own request at the head of LRQ."),
      ("Exit CS", "one FLUSH to the next request in LRQ + REPLY to deferred requests.")]
for i, (h, b) in enumerate(lk):
    numbered(s, Inches(0.6), Inches(2.15) + i * Inches(0.75), Inches(7.2), i + 1, h, b, size=11.5)
rect(s, Inches(8.2), Inches(2.15), Inches(4.6), Inches(4.65), fill=LIGHT)
text(s, Inches(8.4), Inches(2.25), Inches(4.3), Inches(4.5),
     [[("Messages per CS (paper §4.1)", {"bold": True, "size": 13, "color": ORANGE})],
      ["2N − |CSet|  if a lower-priority concurrent request exists"],
      ["2N − 1 − |CSet|  otherwise"],
      ["2(N − 1)  when requests are serialised"],
      [("Our run, N = 3, all concurrent", {"bold": True, "size": 13, "color": ORANGE})],
      ["3 + 3 + 2 = 8 messages (6 REQUEST + 2 FLUSH)"],
      ["vs 12 for Ricart–Agrawala, 18 for Lamport"],
      [("Assumptions", {"bold": True, "size": 13, "color": ORANGE})],
      ["FIFO channels, no crashes - our demo refuses to run with a node offline."]], size=12.5, spacing=1.1, after=3)

# ======================================================= 10 KNAPP + DEADLOCK
s = new_slide("Self-Study", "Knapp's Classification of Distributed Deadlock Detection")
table(s, Inches(0.6), Inches(1.6), Inches(12.1),
      [("Class", "Idea", "In our project"),
       ("Path-pushing", "sites forward wait-for paths until a cycle is visible (Obermarck)", "documented"),
       ("Edge-chasing", "a probe follows wait-for edges; back at the initiator = deadlock (Chandy–Misra–Haas)", "implemented"),
       ("Diffusing computation", "query/reply diffusion over the graph, OR model (CMH)", "documented"),
       ("Global state", "consistent snapshot of all wait-for edges, then search a cycle", "implemented")],
      (2.4, 7.7, 2.0), size=11.5, row_h=0.45)
text(s, Inches(0.6), Inches(4.2), Inches(5.4), Inches(2.7),
     [[("Demo", {"bold": True, "size": 13, "color": ORANGE})],
      ["Each node locks its own unit, then requests the next region's: Pune → Mumbai → Nashik → Pune. No single node can "
       "see the cycle. A probe from Mumbai returns (MUM → NSK → PUNE → MUM); the snapshot method finds the same cycle; "
       "aborting the victim (Pune) breaks it. The lock table is FA-2 only - FA-1 unit data is untouched."]],
     size=12, spacing=1.12)
picture(s, SHOTS / "09a_deadlock_detected.png", Inches(6.2), Inches(4.15), Inches(6.6), Inches(2.55),
        "Real wait-for graph and probe trace")

# ======================================================= 11 OBJECTS + WEB
s = new_slide("Unit IV", "Distributed Objects  ·  Distributed Web-Based Systems")
concept(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(3.05), "Distributed Object-Based System",
        "Objects live on one node and are invoked remotely through an object reference, as in Java RMI / CORBA.",
        ["Reference fa2obj://node/interface/id; naming service in the gateway.",
         "Proxy marshals the call → Invoke RPC → skeleton → TrafficSignal / ResponseUnit servant.",
         "sync_with(ref): a remote reference passed as a parameter; the hosting node calls the other object."])
picture(s, SHOTS / "10_objects.png", Inches(0.6), Inches(4.8), Inches(6.0), Inches(1.95))
concept(s, Inches(6.9), Inches(1.6), Inches(5.9), Inches(3.05), "Distributed Web-Based System",
        "Browser ↔ HTTP/REST + WebSocket ↔ web/app server ↔ back-end services. Caches use validators.",
        ["The dashboard itself is the web client of a multi-tier distributed system.",
         "GET /api/fa2/web/feed returns ETag + Cache-Control: max-age=5.",
         "Revalidation with If-None-Match → 304 Not Modified, empty body (bytes saved)."])
picture(s, SHOTS / "14_web_cache.png", Inches(6.9), Inches(4.8), Inches(5.9), Inches(1.95))

# ======================================================= 12 DFS + SERVERLESS
s = new_slide("Unit IV", "Distributed File System  ·  Serverless Architecture")
concept(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(3.05), "Distributed File System (HDFS / GFS style)",
        "A master keeps metadata; chunk servers keep replicated chunks with checksums.",
        ["Gateway = master: path → chunks → replica nodes. Nodes store 256 B chunks + SHA-256.",
         "2 replicas, least-loaded placement; a read skips offline or corrupted replicas.",
         "Repair re-replicates (also automatically after a node recovers)."])
picture(s, SHOTS / "11_dfs_failover.png", Inches(0.6), Inches(4.8), Inches(6.0), Inches(1.95),
        "Corrupted replica skipped - chunk served by another node")
concept(s, Inches(6.9), Inches(1.6), Inches(5.9), Inches(3.05), "Serverless (Function-as-a-Service)",
        "Developers deploy stateless functions; the platform creates, reuses and destroys instances on demand.",
        ["Local FaaS emulation: each instance is a separate process.",
         "Cold start (new process, ~80-100 ms) vs warm start (same PID, < 1 ms); idle 60 s → scale to zero.",
         "Triggers: INCIDENT_CREATED → classify_incident; RESOLVED → incident_report → stored in the DFS."])
picture(s, SHOTS / "12_faas.png", Inches(6.9), Inches(4.8), Inches(5.9), Inches(1.95), "Cold and warm invocations")

# ======================================================= 13 LEDGER
s = new_slide("Self-Study", "Blockchain / Distributed Ledger: Append-Only Audit Ledger")
items = [("Hash chain", "block = index, time, proposer, previous hash, nonce, events; hash = SHA-256 of all of them."),
         ("Toy proof-of-work", "the proposer searches a nonce so the hash starts with 000 (illustrates mining)."),
         ("Leader ordering, majority commit", "the elected leader proposes; committed when 2 of 3 replicas stored it."),
         ("Replication & catch-up", "each node has its own SQLite replica; a node that was down fetches missing blocks."),
         ("Tamper evidence", "editing an old block breaks its hash link - detected, then repaired from a valid peer."),
         ("Honest limits", "crash-fault tolerant like a permissioned ledger - not Byzantine-fault tolerant.")]
for i, (h, b) in enumerate(items):
    numbered(s, Inches(0.6), Inches(1.65) + i * Inches(0.85), Inches(6.2), i + 1, h, b, size=12)
picture(s, SHOTS / "13a_ledger_tampered.png", Inches(7.1), Inches(1.6), Inches(5.7), Inches(3.4),
        "Nashik replica tampered → broken at block #1")
text(s, Inches(7.1), Inches(5.4), Inches(5.7), Inches(1.4),
     [[("Recorded events: ", {"bold": True, "size": 12, "color": ORANGE}),
       ("incidents, node failures/recoveries, leader elections, CS entries, deadlocks, clock syncs, DFS writes - "
        "batched into a block every 4 s.", {"size": 12})]], spacing=1.15)

# ======================================================= 14 CASE STUDIES
s = new_slide("Unit IV", "Industry Case Studies")
table(s, Inches(0.6), Inches(1.6), Inches(12.1),
      [("System", "What it is", "Key distributed-systems ideas", "Same idea in our project"),
       ("Megaport", "Network-as-a-Service", "Private software-defined backbone; Virtual Cross Connects on demand; Megaport Cloud Router", "Node↔node P2P links as a private regional interconnect"),
       ("Cloudflare", "Global anycast edge network", "Nearest-edge routing, edge caching, Workers (V8 isolates), Durable Objects", "ETag/304 caching, FaaS runtime, single-host remote objects"),
       ("AWS", "Public cloud", "Regions + Availability Zones, replicated S3/DynamoDB, Lambda; Dynamo paper: consistent hashing, vector clocks", "Vector clocks, replicated DFS, majority-commit ledger, FaaS metering"),
       ("Apache Hadoop", "Big-data storage & processing", "HDFS NameNode + DataNodes, 3× block replication, checksums; MapReduce / YARN", "DFS master = NameNode, nodes = DataNodes, re-replication"),
       ("Kubernetes", "Container orchestration", "etcd (Raft) state store, reconciliation controllers, leader election, Services", "Bully/Ring election, beacon health checks, names vs addresses")],
      (1.6, 2.2, 4.6, 3.7), size=11, row_h=0.75)
text(s, Inches(0.6), Inches(6.55), Inches(12.1), Inches(0.4),
     [[("Case studies are documented only - no cloud infrastructure was added to the project.", {"size": 12, "italic": True, "color": GREY})]])

# ======================================================= 15 TESTING
s = new_slide("Verification", "Testing, Results and What Is Simulated")
table(s, Inches(0.6), Inches(1.6), Inches(6.0),
      [("Test", "Result"), ("FA-1 regression (tools/test_system.py)", "40 / 40 passed"),
       ("FA-2 end-to-end (tools/test_fa2.py)", "69 / 69 passed"),
       ("Lamport / vector clock violations", "0"), ("Snapshot invariant (30 vehicles)", "holds"),
       ("Mutex lost updates with LK / RA / Lamport", "0"), ("Leader failure → re-election → recovery", "passed"),
       ("Dashboard click-through (headless Chromium)", "0 console errors")],
      (4.0, 2.0), size=11.5)
table(s, Inches(6.9), Inches(1.6), Inches(5.9),
      [("Real", "Simulated / injected"),
       ("Messages between 3 OS processes over gRPC", "Clock offset and drift per node"),
       ("Lamport/vector clocks, markers, probes, timeouts", "Channel delays in two demos"),
       ("DFS files, checksums, replicas", "Node failure = refuse-all mode"),
       ("FaaS processes, cold/warm starts", "Patrol vehicles, FA-2 lock table"),
       ("Ledger hashes, replication, tamper check", "Case studies (documentation)")],
      (3.0, 2.9), size=11.5)
text(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.3),
     [[("Note: ", {"bold": True, "size": 12, "color": ORANGE}),
       ("one FA-1 check ('traffic density changes over time') is timing-dependent in the original FA-1 code as well "
        "(3 s sensor tick vs 2 s heartbeat vs 3.5 s wait) - measured on the unmodified code and documented, not weakened.",
        {"size": 12, "color": GREY})]], spacing=1.15)

# ======================================================= 16 LIMITATIONS + THANKS
s = new_slide("Conclusion", "Limitations and Conclusion")
lim = [("Reliable-process assumption", "Mutual exclusion and snapshot need all nodes up (demo refuses with a node down)."),
       ("Gateway as single point", "DFS master and ledger service live in the gateway, like an HDFS NameNode without HA."),
       ("Clock steps", "Clocks are stepped, not slewed; skew is simulated on one host."),
       ("No Byzantine tolerance", "The ledger tolerates crashes, not malicious nodes.")]
for i, (h, b) in enumerate(lim):
    x = Inches(0.8) + (i % 2) * Inches(6.2)
    y = Inches(1.8) + (i // 2) * Inches(1.45)
    numbered(s, x, y, Inches(5.6), i + 1, h, b, size=12.5)
rect(s, Inches(0.6), Inches(4.75), Inches(12.1), Inches(0.9), fill=LIGHT)
text(s, Inches(0.8), Inches(4.75), Inches(11.7), Inches(0.9),
     [[("The same three traffic nodes now synchronise clocks, order events, elect a leader, share a resource safely, "
        "detect deadlocks, store replicated files, run serverless functions and keep a tamper-evident audit trail.",
        {"size": 13})]], anchor=MSO_ANCHOR.MIDDLE, spacing=1.15)
rect(s, Inches(0.6), Inches(5.85), Inches(12.1), Inches(1.1), fill=NAVY)
text(s, Inches(0.6), Inches(5.85), Inches(12.1), Inches(1.1),
     [[("Thank You", {"bold": True, "size": 32, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("saved", OUT, "-", len(prs.slides), "slides")
